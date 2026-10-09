"""
Text out of course material, page by page.

PDFs keep their page numbers and slide decks their slide numbers, so a
citation can say "p. 4" or "slide 12". Word documents have no reliable page
numbers until something lays them out, so they come back as one unnumbered
page. Canvas pages, the syllabus, announcements and assignment descriptions
are HTML, reduced here to plain text: nothing from Canvas is ever rendered as
HTML anywhere in the app.

Files are untrusted input. Office files are zip archives, so their unpacked
size is checked before parsing, and every format stops at a page and character
limit rather than reading without end.
"""
import io
import re
import zipfile
from dataclasses import dataclass
from html import unescape
from html.parser import HTMLParser
from pathlib import PurePath
from typing import List, Optional

MAX_PAGES = 500
MAX_CHARS = 2_000_000
MAX_UNZIPPED_BYTES = 300 * 1024 * 1024

PDF = "application/pdf"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

SUPPORTED = {".pdf": PDF, ".pptx": PPTX, ".docx": DOCX}


class UnreadableDocument(Exception):
    """The file could not be read as the format it claims to be."""


@dataclass
class Page:
    number: Optional[int]
    text: str


def kind_of(filename: str, mime_type: Optional[str]) -> Optional[str]:
    """The supported MIME type for a file, by its declared type or else its extension."""
    if mime_type in SUPPORTED.values():
        return mime_type
    return SUPPORTED.get(PurePath(filename or "").suffix.lower())


def _tidy(text: str) -> str:
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\s*\n\s*", "\n", text)
    return text.strip()


def _capped(pages: List[Page]) -> List[Page]:
    kept, used = [], 0
    for page in pages[:MAX_PAGES]:
        text = _tidy(page.text)
        if not text:
            continue
        text = text[: MAX_CHARS - used]
        kept.append(Page(page.number, text))
        used += len(text)
        if used >= MAX_CHARS:
            break
    return kept


def _check_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile:
        raise UnreadableDocument("not a valid Office file")
    if total > MAX_UNZIPPED_BYTES:
        raise UnreadableDocument("unpacks to more than the size limit")


def from_pdf(data: bytes) -> List[Page]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnreadableDocument("the PDF is password protected")
        return _capped([
            Page(number, page.extract_text() or "")
            for number, page in enumerate(reader.pages[:MAX_PAGES], start=1)
        ])
    except UnreadableDocument:
        raise
    except Exception as e:
        raise UnreadableDocument(f"not a readable PDF ({type(e).__name__})")


def from_pptx(data: bytes) -> List[Page]:
    from pptx import Presentation

    _check_zip(data)
    try:
        deck = Presentation(io.BytesIO(data))
        pages = []
        for number, slide in enumerate(deck.slides, start=1):
            if number > MAX_PAGES:
                break
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        parts.append(" | ".join(cell.text for cell in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                parts.append(slide.notes_slide.notes_text_frame.text)
            pages.append(Page(number, "\n".join(parts)))
        return _capped(pages)
    except Exception as e:
        raise UnreadableDocument(f"not a readable slide deck ({type(e).__name__})")


def from_docx(data: bytes) -> List[Page]:
    import docx

    _check_zip(data)
    try:
        document = docx.Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return _capped([Page(None, "\n".join(parts))])
    except Exception as e:
        raise UnreadableDocument(f"not a readable Word document ({type(e).__name__})")


def extract(data: bytes, mime_type: str) -> List[Page]:
    if mime_type == PDF:
        return from_pdf(data)
    if mime_type == PPTX:
        return from_pptx(data)
    if mime_type == DOCX:
        return from_docx(data)
    raise UnreadableDocument(f"unsupported type {mime_type}")


# --- HTML ---------------------------------------------------------------------

_BLOCK = {
    "p", "div", "br", "li", "ul", "ol", "tr", "table", "section", "article",
    "h1", "h2", "h3", "h4", "h5", "h6", "blockquote", "pre", "hr",
}
_SKIP = {"script", "style", "noscript", "template", "iframe", "object", "svg"}


class _TextOnly(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: List[str] = []
        self._skipping = 0

    def handle_starttag(self, tag, attrs):
        if tag in _SKIP:
            self._skipping += 1
        elif tag in _BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in _SKIP and self._skipping:
            self._skipping -= 1
        elif tag in _BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skipping:
            self.parts.append(data)


def html_to_text(html: Optional[str]) -> str:
    if not html:
        return ""
    parser = _TextOnly()
    try:
        parser.feed(html[: MAX_CHARS * 2])
        parser.close()
    except Exception:
        return _tidy(unescape(re.sub(r"<[^>]+>", " ", html)))[:MAX_CHARS]
    return _tidy("".join(parser.parts))[:MAX_CHARS]

"""
A stand-in Canvas for the tests, built on httpx.MockTransport. No test ever
talks to a real Canvas.

API responses are the recorded JSON in tests/fixtures/canvas, named after the
path: /api/v1/courses/101/files is courses_101_files.json. A fixture shaped
{"pages": [[...], [...]]} is served one page at a time with Canvas's Link
header. Files download the way Canvas serves them: the school's host answers
with a redirect to its file store on another host, which sends the bytes.

Every request is recorded, so a test can check what was asked for, with which
headers, and how many times.
"""
import io
import json
from pathlib import Path
from typing import Callable, Dict, List, Optional

import httpx

from api.services import net_guard
from api.services.canvas_client import CanvasClient

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "canvas"
BASE = "https://canvas.test.edu"
FILE_HOST = "files.canvas-cdn.test"
TOKEN = "7~personal-token-for-tests"

# What each test hostname "resolves" to. Real DNS is never used.
ADDRESSES = {
    "canvas.test.edu": ["93.184.216.34"],
    "files.canvas-cdn.test": ["151.101.1.1"],
    "other.test.edu": ["93.184.216.35"],
    "intranet.test": ["10.0.0.5"],
    "metadata.test": ["169.254.169.254"],
    "loopback.test": ["127.0.0.1"],
    "cgnat.test": ["100.64.1.1"],
    "v6-loopback.test": ["::1"],
    "mixed.test": ["93.184.216.36", "192.168.1.10"],
}


def fake_resolve(host: str, port: int) -> List[str]:
    if host not in ADDRESSES:
        raise OSError(f"unknown test host {host}")
    return ADDRESSES[host]


def use_fake_dns() -> None:
    net_guard.resolve = fake_resolve


# --- binary fixtures ------------------------------------------------------------


def make_pdf(pages: List[str]) -> bytes:
    """A small, valid PDF with one line of text per page."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream.decode('latin-1')}\nendstream")
        content_ref = len(objects)
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_ref} 0 R >>"
        )
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_pptx(slides: List[str]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    for text in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = text.split(".")[0]
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
        box.text_frame.text = text
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def make_docx(paragraphs: List[str]) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def default_files() -> Dict[int, bytes]:
    return {
        9001: make_pdf([
            "The first law states that energy is conserved in a closed system.",
            "Work done by a gas equals the integral of pressure over volume.",
        ]),
        9002: make_pptx([
            "Entropy. Entropy measures the number of microstates of a system.",
            "Second law. The entropy of an isolated system never decreases.",
        ]),
        9003: make_docx(["Reading guide: chapter 2 covers heat capacity and calorimetry."]),
    }


# --- the fake ---------------------------------------------------------------------


class FakeCanvas:
    def __init__(self) -> None:
        self.requests: List[httpx.Request] = []
        self.sleeps: List[float] = []
        self.files = default_files()
        # Path -> handler, for a test to replace one endpoint.
        self.overrides: Dict[str, Callable[[httpx.Request], httpx.Response]] = {}
        # Fixture name -> JSON, for a test to change what Canvas returns.
        self.edits: Dict[str, object] = {}
        self.throttle_next = 0

    # -- what tests check -------------------------------------------------------

    def paths(self) -> List[str]:
        return [r.url.path for r in self.requests]

    def downloads(self) -> List[httpx.Request]:
        return [r for r in self.requests if r.url.host == FILE_HOST]

    # -- serving -----------------------------------------------------------------

    def fixture(self, name: str):
        if name in self.edits:
            return self.edits[name]
        path = FIXTURES / f"{name}.json"
        if not path.exists():
            return None
        return json.loads(path.read_text())

    def handle(self, request: httpx.Request) -> httpx.Response:
        self.requests.append(request)
        path = request.url.path

        if self.throttle_next:
            self.throttle_next -= 1
            return httpx.Response(
                403, text="403 Forbidden (Rate Limit Exceeded)", headers={"X-Rate-Limit-Remaining": "0"}
            )
        if path in self.overrides:
            return self.overrides[path](request)

        if request.url.host == FILE_HOST:
            file_id = int(path.rsplit("/", 1)[-1])
            if file_id not in self.files:
                return httpx.Response(404)
            return httpx.Response(200, content=self.files[file_id])

        if request.url.host != "canvas.test.edu":
            return httpx.Response(502, text="unexpected host")

        if path == "/login/oauth2/token":
            if request.method == "DELETE":
                return httpx.Response(200, json={})
            form = dict(httpx.QueryParams(request.content.decode()))
            name = "oauth_refresh" if form.get("grant_type") == "refresh_token" else "oauth_token"
            if form.get("code") == "bad-code":
                return httpx.Response(400, json={"error": "invalid_grant", "error_description": "bad code"})
            return httpx.Response(200, json=self.fixture(name))

        if request.headers.get("Authorization") not in {f"Bearer {t}" for t in self.valid_tokens}:
            return httpx.Response(401, json={"errors": [{"message": "Invalid access token."}]})

        if path.startswith("/files/") and path.endswith("/download"):
            file_id = path.split("/")[2]
            return httpx.Response(302, headers={"Location": f"https://{FILE_HOST}/blobs/{file_id}"})

        name = path.removeprefix("/api/v1/").strip("/").replace("/", "_")
        data = self.fixture(name)
        if data is None:
            return httpx.Response(404, json={"errors": [{"message": "The specified resource does not exist."}]})

        headers = {"X-Rate-Limit-Remaining": "650.0"}
        if isinstance(data, dict) and "pages" in data:
            page = int(request.url.params.get("page", "1"))
            body = data["pages"][page - 1]
            if page < len(data["pages"]):
                headers["Link"] = (
                    f'<{BASE}{path}?page={page + 1}&per_page=100>; rel="next", '
                    f'<{BASE}{path}?page=1&per_page=100>; rel="first"'
                )
            return httpx.Response(200, json=body, headers=headers)
        return httpx.Response(200, json=data, headers=headers)

    valid_tokens = {TOKEN, "1~oauth-access-token-from-canvas", "1~refreshed-access-token"}

    @property
    def transport(self) -> httpx.MockTransport:
        return httpx.MockTransport(self.handle)

    def http(self) -> httpx.Client:
        return httpx.Client(transport=self.transport, follow_redirects=False)

    def client(self, base_url: str, token: str) -> CanvasClient:
        return CanvasClient(base_url, token, http=self.http(), sleep=self.sleeps.append)


def install(fake: Optional[FakeCanvas] = None) -> FakeCanvas:
    """Route every Canvas call the app makes to `fake`."""
    from api.services import canvas_auth, canvas_sync

    fake = fake or FakeCanvas()
    use_fake_dns()
    canvas_sync.client_factory = fake.client
    canvas_auth.http_factory = fake.http
    return fake
